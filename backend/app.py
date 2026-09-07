"""API de publicaciones y cuadros estadísticos para Python/MySQL."""
from functools import wraps
import os
from pathlib import Path

import pymysql
from dotenv import load_dotenv
from flask import Flask, jsonify, request
from openpyxl import load_workbook

load_dotenv()
app = Flask(__name__)
DATA_ROOT = Path(os.getenv("DATA_ROOT", "datos_OE")).resolve()


def db_connection():
    """Abre una conexión MySQL usando variables de entorno del despliegue."""
    return pymysql.connect(
        host=os.environ["MYSQL_HOST"],
        port=int(os.getenv("MYSQL_PORT", "3306")),
        user=os.environ["MYSQL_USER"],
        password=os.environ["MYSQL_PASSWORD"],
        database=os.environ["MYSQL_DATABASE"],
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=True,
        connect_timeout=int(os.getenv("MYSQL_CONNECT_TIMEOUT", "10")),
    )


def api_errors(view):
    """Convierte errores de conexión y consultas en respuestas JSON consistentes."""
    @wraps(view)
    def wrapped(*args, **kwargs):
        try:
            return view(*args, **kwargs)
        except pymysql.MySQLError:
            app.logger.exception("Error de base de datos")
            return jsonify(error="No fue posible consultar la base de datos"), 503
        except OSError:
            app.logger.exception("Error leyendo archivo de datos")
            return jsonify(error="No fue posible leer el cuadro estadístico"), 503

    return wrapped


@app.after_request
def add_cors_headers(response):
    """Permite consumir la API desde el frontend institucional configurado."""
    response.headers["Access-Control-Allow-Origin"] = os.getenv("CORS_ALLOWED_ORIGINS", "*")
    response.headers["Access-Control-Allow-Headers"] = "Content-Type"
    response.headers["Access-Control-Allow-Methods"] = "GET, OPTIONS"
    return response


@app.get("/api/health")
def health():
    """Comprueba que el proceso Flask está disponible."""
    return jsonify(status="ok", service="demografia-de-empresas-api")


@app.get("/api/categorias")
@api_errors
def categorias():
    """Lista las categorías que tienen al menos un recurso publicado."""
    with db_connection() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT c.slug, c.nombre, c.seccion, c.orden,
                   COUNT(r.id) AS recursos_publicados
              FROM categorias AS c
              JOIN recursos AS r ON r.categoria_id = c.id
             WHERE r.estado = 'publicado'
             GROUP BY c.id, c.slug, c.nombre, c.seccion, c.orden
             ORDER BY c.orden, c.nombre
            """
        )
        return jsonify(cursor.fetchall())


def _published_resource(filename):
    """Busca el archivo por nombre y confirma que esté publicado en el catálogo."""
    with db_connection() as connection, connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT nombre_archivo, ruta_relativa, tipo_archivo
              FROM recursos
             WHERE nombre_archivo = %s AND estado = 'publicado'
             LIMIT 1
            """,
            (filename,),
        )
        return cursor.fetchone()


@app.get("/api/cuadros/<path:filename>")
@api_errors
def cuadro(filename):
    """Devuelve las filas de un XLSX publicado, con paginación opcional."""
    if Path(filename).name != filename or not filename.lower().endswith(".xlsx"):
        return jsonify(error="Sólo se permiten nombres de archivos XLSX"), 400

    resource = _published_resource(filename)
    if not resource:
        return jsonify(error="Cuadro no encontrado o no publicado"), 404

    relative_path = Path(resource["ruta_relativa"])
    file_path = (DATA_ROOT / relative_path.relative_to("datos_OE")).resolve()
    if DATA_ROOT not in file_path.parents or not file_path.is_file():
        return jsonify(error="El archivo publicado no está disponible"), 404

    try:
        offset = max(int(request.args.get("offset", "0")), 0)
        limit = min(max(int(request.args.get("limit", "10000")), 1), 50000)
    except ValueError:
        return jsonify(error="offset y limit deben ser números enteros"), 400

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else f"columna_{i + 1}"
               for i, value in enumerate(next(rows, ()))]
    data = []
    for index, values in enumerate(rows):
        if index < offset:
            continue
        if len(data) >= limit:
            break
        data.append({
            headers[i]: value
            for i, value in enumerate(values)
            if i < len(headers)
        })
    workbook.close()

    return jsonify(
        archivo=resource["nombre_archivo"],
        hoja=sheet.title,
        offset=offset,
        limit=limit,
        filas=len(data),
        datos=data,
    )


@app.get("/api/publicaciones")
@api_errors
def publicaciones():
    """Entrega recursos publicados con filtros opcionales."""
    conditions = ["r.estado = 'publicado'"]
    params = []

    categoria = request.args.get("categoria", "").strip()
    tipo = request.args.get("tipo", "").strip().upper()
    query = request.args.get("q", "").strip()
    incluir_documentacion = request.args.get("documentacion", "0") == "1"

    if categoria:
        conditions.append("c.slug = %s")
        params.append(categoria)
    if tipo:
        conditions.append("UPPER(r.tipo_archivo) = %s")
        params.append(tipo)
    if query:
        conditions.append("(r.titulo_publico LIKE %s OR r.nombre_archivo LIKE %s)")
        pattern = f"%{query}%"
        params.extend([pattern, pattern])
    if not incluir_documentacion:
        conditions.append("c.seccion = 'publicaciones'")

    sql = f"""
        SELECT c.slug AS categoria, c.nombre AS categoria_nombre,
               c.seccion, r.titulo_publico, r.nombre_archivo,
               r.ruta_relativa, r.tipo_archivo, r.fecha_publicacion, r.orden
          FROM recursos AS r
          JOIN categorias AS c ON c.id = r.categoria_id
         WHERE {" AND ".join(conditions)}
         ORDER BY c.orden, r.orden, r.titulo_publico
    """

    with db_connection() as connection, connection.cursor() as cursor:
        cursor.execute(sql, params)
        return jsonify(cursor.fetchall())


@app.errorhandler(404)
def not_found(_error):
    """Mantiene el formato JSON para rutas inexistentes."""
    return jsonify(error="Ruta no encontrada"), 404


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
